"""Gera dados/importacao.sql com o perfil do Guilherme e as visitas aprovadas, já com os
códigos oficiais de dados/lugares-codigos.xlsx. O SQL é rodado à mão no SQL Editor do Supabase.

Ele só importa se o usuário ainda não tiver nenhuma visita, para não duplicar nada.

Uso: python scripts/gerar_importacao.py <User UID do Supabase>
"""
import os
import sys

import openpyxl

RAIZ = os.path.join(os.path.dirname(__file__), "..")
VISITAS = os.path.join(RAIZ, "dados", "visitas-aprovadas.xlsx")
CODIGOS = os.path.join(RAIZ, "dados", "lugares-codigos.xlsx")
DESTINO = os.path.join(RAIZ, "dados", "importacao.sql")


def texto(valor):
    return "null" if valor in (None, "") else "'" + str(valor).replace("'", "''") + "'"


def main():
    uid = sys.argv[1]
    codigos = {}
    for pais, pais_cod, n1, n1_cod, condado, condado_cod, mun, mun_cod, oficial, _, situacao in \
            openpyxl.load_workbook(CODIGOS).active.iter_rows(min_row=2, values_only=True):
        if not mun_cod:
            raise SystemExit(f"Sem código: {pais} / {mun} ({situacao})")
        codigos[(pais, n1, mun)] = (pais_cod, n1_cod, condado, condado_cod, mun_cod)

    linhas = []
    for ano, pais, n1, _, mun, local, obs in openpyxl.load_workbook(VISITAS).active.iter_rows(min_row=2, values_only=True):
        pais_cod, n1_cod, condado, condado_cod, mun_cod = codigos[(pais, n1, mun)]
        valores = [texto(uid), str(int(ano)), texto(pais_cod), texto(n1_cod), texto(n1),
                   texto(condado_cod), texto(condado), texto(mun_cod), texto(mun), texto(local), texto(obs)]
        linhas.append("  (" + ", ".join(valores) + ")")

    sql = f"""-- Órbita: importação das visitas da planilha (gerado por scripts/gerar_importacao.py)
-- Roda inteiro ou não roda nada. Recusa se já houver visitas, para não duplicar.

begin;

insert into public.perfis (id, usuario, nome, publico)
values ({texto(uid)}, 'guilherme', 'Guilherme Lopes', true)
on conflict (id) do nothing;

do $$
begin
  if exists (select 1 from public.visitas where usuario_id = {texto(uid)}) then
    raise exception 'Já existem visitas para este usuário; importação cancelada para não duplicar.';
  end if;
end $$;

insert into public.visitas
  (usuario_id, ano, pais_codigo, nivel1_codigo, nivel1_nome, condado_codigo, condado_nome,
   nivel2_codigo, nivel2_nome, local_nome, observacao)
values
{",\n".join(linhas)};

commit;

select count(*) as visitas_importadas from public.visitas where usuario_id = {texto(uid)};
"""
    with open(DESTINO, "w", encoding="utf-8") as f:
        f.write(sql)
    print(f"{DESTINO}: {len(linhas)} visitas")


if __name__ == "__main__":
    main()
