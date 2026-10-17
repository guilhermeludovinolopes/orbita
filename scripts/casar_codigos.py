"""Casa cada lugar de dados/visitas-aprovadas.xlsx com o código oficial do seu país e gera
dados/lugares-codigos.xlsx para conferência.

Fontes (baixadas para uma pasta passada como argumento):
  Brasil   IBGE, API de localidades (ibge-municipios.json)
  EUA      Census Bureau, gazetteer 2024 (census/2024_Gaz_counties_national.txt e _place_)
  Itália   ISTAT, Elenco comuni (istat-comuni.csv) + ISO 3166-2:IT (iso-3166-2-it.txt)
  Suíça    BFS, registro de comunas (bfs-comunas.csv)
  Portugal INE, códigos DICO via json.geoapi.pt (consulta por nome)

Uso: python scripts/casar_codigos.py <pasta das fontes>
"""
import csv
import json
import os
import re
import sys
import unicodedata
import urllib.parse
import urllib.request

import openpyxl
from openpyxl.styles import Font, PatternFill

RAIZ = os.path.join(os.path.dirname(__file__), "..")
ORIGEM = os.path.join(RAIZ, "dados", "visitas-aprovadas.xlsx")
DESTINO = os.path.join(RAIZ, "dados", "lugares-codigos.xlsx")
FONTES = sys.argv[1]

UFS = {"Acre": "AC", "Ceará": "CE", "Distrito Federal": "DF", "Goiás": "GO", "Mato Grosso do Sul": "MS",
       "Minas Gerais": "MG", "Pará": "PA", "Piauí": "PI", "Rio de Janeiro": "RJ", "Rondônia": "RO",
       "Santa Catarina": "SC", "São Paulo": "SP", "Tocantins": "TO"}
ESTADOS_EUA = {"California": "CA", "Florida": "FL", "Nevada": "NV"}
CANTOES = {"Graubünden": "GR", "Ticino": "TI"}


def normalizar(texto):
    sem_acento = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sem_acento.lower())


def carregar_brasil():
    with open(os.path.join(FONTES, "ibge-municipios.json"), encoding="utf-8") as f:
        municipios = json.load(f)
    indice = {}
    for m in municipios:
        uf = (m.get("microrregiao") or {}).get("mesorregiao", {}).get("UF", {}).get("sigla") \
            or m["regiao-imediata"]["regiao-intermediaria"]["UF"]["sigla"]
        indice[(uf, normalizar(m["nome"]))] = (str(m["id"]), m["nome"])
    return indice


def carregar_census(arquivo):
    with open(os.path.join(FONTES, "census", arquivo), encoding="latin-1") as f:
        leitor = csv.DictReader(f, delimiter="\t")
        return [{k.strip(): v.strip() for k, v in linha.items()} for linha in leitor]


def carregar_italia():
    with open(os.path.join(FONTES, "istat-comuni.csv"), encoding="latin-1", newline="") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    cab = [c.replace("\n", " ") for c in linhas[0]]
    col = lambda nome: next(i for i, c in enumerate(cab) if c.startswith(nome))
    i_cod, i_nome, i_reg = col("Codice Comune formato alfanumerico"), col("Denominazione in italiano"), col("Denominazione Regione")
    # Província (ou cidade metropolitana): o nível entre a região e o comune
    i_prov_cod = col("Codice dell'Unità territoriale sovracomunale")
    i_prov_nome = col("Denominazione dell'Unità territoriale sovracomunale")
    comuni = {
        (normalizar(l[i_reg]), normalizar(l[i_nome])): (l[i_cod], l[i_nome], l[i_prov_cod], l[i_prov_nome])
        for l in linhas[1:] if len(l) > i_reg
    }
    with open(os.path.join(FONTES, "iso-3166-2-it.txt"), encoding="utf-8") as f:
        iso = {normalizar(nome): cod for cod, nome in re.findall(r"\{\{mono\|(IT-\d{2})\}\}\s*\n\|\s*\{\{flag\|([^}|]+)", f.read())}
    return comuni, iso


def carregar_suica():
    with open(os.path.join(FONTES, "bfs-comunas.csv"), encoding="utf-8") as f:
        linhas = list(csv.DictReader(f))
    # Distrito (ou região): o nível entre o cantão e a comuna; a comuna aponta para ele em "Parent"
    distritos = {l["HistoricalCode"]: (l["BfsCode"], l["Name"]) for l in linhas if l["Level"] == "2"}
    return {normalizar(l["Name"]): (l["BfsCode"], l["Name"], *distritos[l["Parent"]]) for l in linhas if l["Level"] == "3"}


def buscar_portugal(nome):
    url = "https://json.geoapi.pt/municipio/" + urllib.parse.quote(nome)
    with urllib.request.urlopen(url, timeout=30) as r:
        dados = json.loads(r.read().decode("utf-8"))
    return dados.get("codigoine"), dados.get("nome")


def main():
    brasil = carregar_brasil()
    condados = carregar_census("2024_Gaz_counties_national.txt")
    cidades = carregar_census("2024_Gaz_place_national.txt")
    comuni, iso_it = carregar_italia()
    suica = carregar_suica()

    ws = openpyxl.load_workbook(ORIGEM).active
    lugares = sorted({(p, n1, c, m) for _, p, n1, c, m, *_ in ws.iter_rows(min_row=2, values_only=True)},
                     key=lambda x: (x[0], x[1], x[2] or "", x[3]))

    saida = []
    for pais, n1, condado, mun in lugares:
        r = {"pais": pais, "pais_cod": "", "n1": n1, "n1_cod": "", "condado": condado or "", "condado_cod": "",
             "mun": mun, "mun_cod": "", "mun_oficial": "", "fonte": "", "situacao": "OK"}
        try:
            if pais == "Brasil":
                uf = UFS[n1]
                r.update(pais_cod="BR", n1_cod=f"BR-{uf}", fonte="IBGE")
                r["mun_cod"], r["mun_oficial"] = brasil[(uf, normalizar(mun))]
            elif pais == "Estados Unidos":
                uf = ESTADOS_EUA[n1]
                r.update(pais_cod="US", n1_cod=f"US-{uf}", fonte="Census 2024")
                cond = next(c for c in condados if c["USPS"] == uf and c["NAME"] == condado)
                r["condado_cod"] = cond["GEOID"]
                alvo = normalizar(mun)
                cand = [c for c in cidades if c["USPS"] == uf and normalizar(re.sub(r" (city|town|village|CDP)$", "", c["NAME"])) == alvo]
                cand.sort(key=lambda c: not c["NAME"].endswith(" city"))  # prefere a cidade oficial ao CDP
                r["mun_cod"], r["mun_oficial"] = cand[0]["GEOID"], cand[0]["NAME"]
            elif pais == "Itália":
                r.update(pais_cod="IT", n1_cod=iso_it[normalizar(n1)], fonte="ISTAT + ISO 3166-2")
                r["mun_cod"], r["mun_oficial"], r["condado_cod"], r["condado"] = comuni[(normalizar(n1), normalizar(mun))]
            elif pais == "Suíça":
                r.update(pais_cod="CH", n1_cod=f"CH-{CANTOES[n1]}", fonte="BFS")
                r["mun_cod"], r["mun_oficial"], r["condado_cod"], r["condado"] = suica[normalizar(mun)]
            elif pais == "Inglaterra":
                # Decisão do Guilherme: Londres é um lugar só (os 33 distritos juntos), com o
                # mesmo código ONS da região Greater London nos dois níveis
                r.update(pais_cod="GB-ENG", n1_cod="E12000007", mun_cod="E12000007", mun_oficial="Londres",
                         fonte="ONS", situacao="OK")
            elif pais == "Irlanda":
                # Decisão do Guilherme: Dublin City (a prefeitura da cidade, uma das 4 do condado).
                # Código provisório até montarmos o mapa da Irlanda com a fonte oficial.
                r.update(pais_cod="IE", n1_cod="IE-L", condado_cod="IE-D", mun_cod="IE-D-DUBLIN-CITY",
                         mun_oficial="Dublin City", fonte="ISO 3166-2 (cidade provisória)", situacao="CÓDIGO PROVISÓRIO")
            elif pais == "Portugal":
                codigo, oficial = buscar_portugal(mun)
                r.update(pais_cod="PT", n1_cod=f"PT-{codigo[:2]}", mun_cod=codigo, mun_oficial=oficial, fonte="INE (geoapi.pt)")
            else:
                r.update(situacao="DECIDIR", fonte="—")
        except (KeyError, StopIteration, IndexError, TypeError) as erro:
            r.update(situacao=f"NÃO ACHEI ({type(erro).__name__})")
        if r["situacao"] == "OK" and pais != "Inglaterra" and normalizar(r["mun_oficial"]) not in (normalizar(mun), normalizar(mun) + "city"):
            r["situacao"] = "CONFERIR NOME"
        saida.append(r)

    wb = openpyxl.Workbook()
    out = wb.active
    out.title = "Lugares e códigos"
    cab = ["País", "Código país", "Estado / 1º nível", "Código 1º nível", "Condado", "Código condado",
           "Município", "Código município", "Nome oficial na fonte", "Fonte", "Situação"]
    out.append(cab)
    for r in saida:
        out.append([r[k] for k in ["pais", "pais_cod", "n1", "n1_cod", "condado", "condado_cod", "mun", "mun_cod",
                                   "mun_oficial", "fonte", "situacao"]])
    for c in out[1]:
        c.font = Font(bold=True, color="F5F0E8")
        c.fill = PatternFill("solid", fgColor="13294B")
    for linha in out.iter_rows(min_row=2):
        if linha[10].value != "OK":
            for c in linha:
                c.fill = PatternFill("solid", fgColor="E8C9CD")
    for col, largura in zip("ABCDEFGHIJK", [15, 10, 20, 12, 22, 12, 24, 14, 28, 20, 22]):
        out.column_dimensions[col].width = largura
    out.freeze_panes = "A2"
    wb.save(DESTINO)

    for r in saida:
        if r["situacao"] != "OK":
            print(r["situacao"], "|", r["pais"], r["n1"], r["condado"], r["mun"], "->", r["mun_oficial"], r["mun_cod"])
    print(len(saida), "lugares;", sum(r["situacao"] == "OK" for r in saida), "OK")


if __name__ == "__main__":
    main()
