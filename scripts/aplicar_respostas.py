"""Aplica as respostas do Guilherme à tabela de correções e gera a lista final de visitas
aprovadas, com o condado (EUA, Irlanda) numa coluna própria entre o estado e o município."""
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill

RESPOSTAS = os.path.expanduser("~") + "/Downloads/Órbita - Correções para aprovar.xlsx"
DESTINO = os.path.join(os.path.dirname(__file__), "..", "dados", "visitas-aprovadas.xlsx")

# Países onde existe um nível de condado entre o estado e o município
PAISES_COM_CONDADO = {"Estados Unidos", "Irlanda"}


def main():
    ws = openpyxl.load_workbook(RESPOSTAS, data_only=True).worksheets[0]
    visitas = []
    for ano, _, _, cidade_antes, pais, nivel1, unidade, local, obs, *_ in ws.iter_rows(min_row=2, values_only=True):
        condado = None
        municipio = unidade
        if pais in PAISES_COM_CONDADO:
            # Na tabela anterior o condado estava no lugar do município e a cidade virou "local"
            condado, municipio = unidade, local
            local = None
        visitas.append([ano, pais, nivel1, condado, municipio, local, obs])

    wb = openpyxl.Workbook()
    out = wb.active
    out.title = "Visitas aprovadas"
    out.append(["Ano", "País", "Estado / 1º nível", "Condado (só onde existe)",
                "Município", "Local específico", "Observação"])
    for v in sorted(visitas, key=lambda v: (v[0], v[1], v[2] or "", v[4] or "")):
        out.append(v)
    for c in out[1]:
        c.font = Font(bold=True, color="F5F0E8")
        c.fill = PatternFill("solid", fgColor="13294B")
    for col, largura in zip("ABCDEFG", [7, 16, 20, 24, 26, 26, 30]):
        out.column_dimensions[col].width = largura
    out.freeze_panes = "A2"
    wb.save(DESTINO)

    com_condado = sum(1 for v in visitas if v[3])
    print(f"{len(visitas)} visitas; {com_condado} com condado")
    for v in visitas:
        if v[3]:
            print(v)


if __name__ == "__main__":
    sys.exit(main())
